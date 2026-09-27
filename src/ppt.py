from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# Helper function
def add_title_slide(title, subtitle):
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.title
    subtitle_shape = slide.placeholders[1]
    title_shape.text = title
    subtitle_shape.text = subtitle
    return slide

def add_content_slide(title, content_list):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.title
    title_shape.text = title
    
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    tf.clear()
    
    for point in content_list:
        p = tf.add_paragraph()
        p.text = point
        p.level = 0
    return slide

# ==================== SLIDE 1: TITLE ====================
add_title_slide(
    "PROJECT WORK PHASE – I\nFIRST REVIEW",
    "IoT Intrusion Detection System using Machine Learning\nDate: 14 July 2026"
)

# ==================== SLIDE 2: PROBLEM STATEMENT ====================
add_content_slide("Problem Statement", [
    "• Traditional signature-based IDS fail against zero-day attacks",
    "• IoT devices have severe resource constraints (low power, memory)",
    "• High risk in smart homes, industries, and critical infrastructure",
    "• Need for lightweight, intelligent, real-time detection system"
])

# ==================== SLIDE 3: OBJECTIVES ====================
add_content_slide("Objectives", [
    "Primary Objective: Develop ML-based NIDS using Random Forest",
    "• Preprocess NSL-KDD dataset",
    "• Train high-accuracy model (>95%)",
    "• Build interactive Streamlit web dashboard",
    "• Implement PCAP file analysis feature",
    "• Achieve real-time threat detection"
])

# ==================== SLIDE 4: CURRENT PROGRESS ====================
add_content_slide("Current Progress (First Review)", [
    "✅ Dataset loaded & preprocessed (NSL-KDD)",
    "✅ Random Forest model trained with 99.80% accuracy",
    "✅ Streamlit Web Dashboard completed (Manual Input)",
    "✅ PCAP file upload & analysis feature implemented",
    "✅ Professional project structure & documentation"
])

# ==================== SLIDE 5: RESULTS ====================
add_content_slide("Results & Achievements", [
    "• Model Accuracy: 99.80%",
    "• Binary Classification (Normal vs Attack)",
    "• Real-time predictions with confidence score",
    "• User-friendly Streamlit interface",
    "• Basic flow-based PCAP analysis"
])

# ==================== SLIDE 6: METHODOLOGY ====================
add_content_slide("Proposed Methodology", [
    "• Data Source: NSL-KDD Dataset",
    "• Preprocessing: Label Encoding + Standard Scaling",
    "• Algorithm: Random Forest Classifier",
    "• Frontend: Streamlit Dashboard",
    "• Features: Duration, src_bytes, dst_bytes, count, serror_rate, TCP Flags etc."
])

# ==================== SLIDE 7: FUTURE WORK ====================
add_content_slide("Challenges & Future Work", [
    "• Improve PCAP feature extraction for better attack detection",
    "• Integrate with modern IoT datasets (BoT-IoT, TON_IoT)",
    "• Add live packet capture using Scapy",
    "• Feature importance visualization",
    "• Multi-class attack type prediction"
])

# ==================== SLIDE 8: TIMELINE ====================
add_content_slide("Timeline", [
    "Week 1-2: Data Preprocessing & Model Training → Completed",
    "Week 3: Streamlit Dashboard & PCAP Feature → Completed",
    "Week 4: Testing, Documentation & Presentation"
])

# ==================== SLIDE 9: REFERENCES ====================
add_content_slide("References", [
    "• NSL-KDD Dataset",
    "• IEEE Papers on Random Forest IDS",
    "• Scikit-learn, Streamlit, Scapy documentation"
])

# Save the file
prs.save("First_Review_IoT_IDS.pptx")
print("✅ First Review PPTX generated successfully!")
print("File saved as: First_Review_IoT_IDS.pptx")
